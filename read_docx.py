import docx
from uyghur.conversion import uey2uly
import re

doc = docx.Document("1_abljzhaoji20211219-维汉句子数据库（阿拉伯数字修改版）(1)(1).docx")
uey = doc.tables[0].cell(0, 1).text
lines = uey2uly(uey).splitlines()
with open("orig.txt", "w") as f:
    f.write(uey)
def write_to(text: str, path):
    with open(path, "w") as f:
        for i, line in enumerate(text.splitlines()):
            line = re.sub(rf"^\d+", "", line)
            line = re.sub(rf"\d+$", "", line)
            line = re.sub(f"{i + 1}", "", line.strip())
            line = line.replace(".", "")
            for symbol in ("?", "؟", "!"):
                if symbol in line:
                    line = line.replace(symbol, "") + symbol
            if line[-1] not in ("?", "؟", "!"):
                line += "."
            line = re.sub(r"\s+", " ", line)
            f.write(line.strip()+"\n")
write_to(uey, "uey.txt")
write_to(uey2uly(uey), "uly.txt")
